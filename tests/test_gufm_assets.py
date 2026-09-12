"""GUFM leftovers: JAXA palette table, Seki-lab blank deck, zone CSV schema."""

from __future__ import annotations

from pathlib import Path

from qgis_mcp_workflows.helpers import (
    SEKILAB_SLIDE_HEIGHT,
    SEKILAB_SLIDE_WIDTH,
    bundled_asset,
)
from qgis_mcp_workflows_plugin.colormaps import JAXA_LULC_CLASSES


def test_jaxa_lulc_table_has_16_classes_including_nodata():
    values = [row[0] for row in JAXA_LULC_CLASSES]
    assert values == list(range(16))
    assert JAXA_LULC_CLASSES[0][2] == "unclassified"
    assert JAXA_LULC_CLASSES[2][2] == "built-up"
    # Built-up is red in the published JAXA legend / GeoTIFF colour table.
    assert JAXA_LULC_CLASSES[2][1][:3] == (255, 0, 0)


def test_gufm_ramp_is_vendored():
    from qgis_mcp_workflows_plugin.colormaps import COLORMAPS

    assert "gufm" in COLORMAPS
    assert COLORMAPS["gufm"][0][1] == (244, 246, 251)
    assert COLORMAPS["gufm"][-1][1] == (184, 92, 0)


def test_sekilab_blank_matches_lab_slide_size():
    path = bundled_asset("assets", "sekilab_blank.pptx")
    assert path, "run: uv run --no-sync --extra pptx scripts/make_sekilab_blank.py"
    from pptx import Presentation

    prs = Presentation(path)
    assert int(prs.slide_width) == SEKILAB_SLIDE_WIDTH
    assert int(prs.slide_height) == SEKILAB_SLIDE_HEIGHT
    assert len(prs.slides) == 0
    assert len(prs.slide_layouts) >= 7


def test_home_counts_csv_uses_n03_zone_id():
    path = bundled_asset("assets", "tokyo23_home_counts.csv")
    assert path, "run: uv run --no-sync --extra drm scripts/build_gufm_zones.py"
    text = Path(path).read_text(encoding="utf-8")
    header, *rows = text.strip().splitlines()
    assert header == "zone_id,n_persons"
    ids = [line.split(",")[0] for line in rows]
    assert ids == [str(z) for z in range(13101, 13124)]


def test_figures_to_pptx_defaults_to_sekilab_size(tmp_path):
    from PIL import Image

    from qgis_mcp_workflows.server import qgis_figures_to_pptx

    img = tmp_path / "f.png"
    Image.new("RGB", (80, 60), color=(20, 80, 120)).save(img)
    out = tmp_path / "out.pptx"
    result = qgis_figures_to_pptx(figure_paths=[str(img)], pptx_path=str(out))
    from pptx import Presentation

    prs = Presentation(result.pptx_path)
    assert int(prs.slide_width) == SEKILAB_SLIDE_WIDTH
    assert result.n_slides_added == 1
